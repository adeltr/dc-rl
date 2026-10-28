"""Pont Excel : non-régression contre le classeur original, règles d'écriture, bout en bout."""
import platform
import re
import zipfile
from datetime import datetime

import openpyxl
import pytest

from bridge import bridge as B
from bridge.engines import EngineUnavailable, select_engine
from bridge.mapping import CAPEX_KEYS, MarketAssumptions, original_market, workbook_inputs
from bridge.spec import INPUTS, SCALAR_OUTPUTS, SERIES_OUTPUTS
from twin.engine import simulate
from twin.params import TwinParams

TOLERANCE = 1e-6
pytestmark = pytest.mark.skipif(not B.SOURCE_WORKBOOK.exists(), reason="classeur absent (model/ non versionné)")

# Sorties qui renvoient directement une cellule du classeur original
DIRECT = {name: re.sub(r"\$", "", formula.split("!")[1]) for name, _d, _u, formula in SCALAR_OUTPUTS
          if formula.startswith("Underwriting!$") and ":" not in formula}
NAMED = {"out_capex_total": "K76", "out_equity": "K79", "out_debt": "K80", "out_ltc_actual": "F80",
         "out_ltc_target": "C80", "out_value_untrended": "I207", "out_value_sale": "K207", "out_it_load_kw": "E131",
         "out_pue": "K23", "out_operation_begin_month": "D113", "out_stabilization_month": "D114",
         "out_analysis_end_month": "K25"}


def close(a, b):
    return abs(a - b) <= TOLERANCE * max(1.0, abs(a))


@pytest.fixture(scope="module")
def original():
    return openpyxl.load_workbook(B.SOURCE_WORKBOOK, data_only=True)


@pytest.fixture(scope="module")
def bridge(tmp_path_factory):
    workbook = tmp_path_factory.mktemp("bridge") / "twin_model.xlsm"
    b = B.ExcelBridge(workbook=workbook)
    yield b
    b.close()


@pytest.fixture(scope="module")
def base_run(bridge, tmp_path_factory):
    export = tmp_path_factory.mktemp("export") / "base.xlsx"
    return bridge.run_inputs({}, export_path=export)


@pytest.fixture(scope="module")
def twin_base():
    return simulate(TwinParams(site_id="us_ny_jfk"))


# ------------------------------------------------------------------ non-régression

def test_every_original_cell_is_reproduced(original, base_run):
    recalculated = openpyxl.load_workbook(base_run.export_path, data_only=True)
    checked, failures = 0, []
    for ws in original.worksheets:
        other = recalculated[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                a = cell.value
                if not isinstance(a, (int, float)) or isinstance(a, bool):
                    continue
                b = other[cell.coordinate].value
                if isinstance(b, datetime) or not isinstance(b, (int, float)) or not close(a, b):
                    failures.append((ws.title, cell.coordinate, a, b))
                checked += 1
    assert checked > 25_000
    assert not failures, failures[:20]


def test_outputs_match_original(original, base_run):
    uw = original["Underwriting"]
    for name, ref in {**DIRECT, **NAMED}.items():
        assert close(uw[ref].value, base_run.outputs[name]), name
    assert base_run.recalc_ok and base_run.solver_iterations == 0

    # sorties nouvelles, recalculées en Python depuis les valeurs en cache du classeur original
    def monthly(row):
        return [uw.cell(row=row, column=19 + m).value for m in range(133)]

    begin = int(uw["D113"].value)
    assert close(sum(monthly(153)[begin:begin + 12]), base_run.outputs["out_opex_energy_y1"])
    years = [uw.cell(row=28, column=19 + m).value for m in range(133)]
    for name, _d, row in SERIES_OUTPUTS:
        values = monthly(row)
        for i, year in enumerate(base_run.years):
            expected = sum(v for v, y in zip(values, years) if y == year and isinstance(v, (int, float)))
            assert close(expected, base_run.series[name][i]), (name, year)

    end = int(uw["K25"].value)
    start = 1 if uw["S76"].value == 0 else 0  # même fenêtre que les XIRR de J229
    flows, dates = monthly(229)[start:end + 1], monthly(213)[start:end + 1]
    rate, t0 = 0.08, dates[0]
    npv = sum(cf / (1 + rate) ** ((d - t0).days / 365) for cf, d in zip(flows, dates))
    assert close(npv, base_run.outputs["out_npv"])


def test_original_market_gives_original_values(bridge, twin_base):
    inputs = workbook_inputs(twin_base, twin_base, original_market(bridge.manifest), bridge.manifest)
    for spec in INPUTS:
        if not spec.name.startswith("tw_memo_"):
            assert inputs[spec.name] == bridge.manifest["inputs"][spec.name]["value"], spec.name


def test_default_base_uses_simulated_energy(bridge, twin_base):
    """Cas de base par défaut : le classeur reçoit le PUE et le coût d'énergie simulés ; l'énergie suit
    l'occupation au-delà de la part fixe simulée ; la taxe foncière suit les coûts durs."""
    run = bridge.run(twin_base, twin_base, MarketAssumptions())
    assert run.inputs["tw_pue"] == pytest.approx(twin_base.pue_annual)
    assert run.outputs["out_utilities_proforma"] == pytest.approx(twin_base.energy_cost_usd)
    assert run.inputs["tw_utilities_fixed"] == pytest.approx(twin_base.fixed_energy_share)
    original = bridge.run_inputs({})
    assert run.inputs["tw_property_tax"] == pytest.approx(original.inputs["tw_property_tax"])  # mêmes coûts durs
    # montée en occupation : la 1re année d'exploitation coûte bien moins qu'avec 100 % fixe
    assert run.outputs["out_opex_energy_y1"] < 0.7 * original.outputs["out_opex_energy_y1"]
    i3 = run.years.index(3)
    assert run.series["out_cfads"][i3] > original.series["out_cfads"][i3]
    # sans dette permanente, les intérêts du prêt de construction forment le service de la dette
    assert run.series["out_debt_service_total"][i3] == pytest.approx(run.series["out_construction_interest_y"][i3])
    # DSCR testé à partir de la première année stabilisée (stabilisation au mois 46 : A5)
    assert run.series["out_dscr"][i3] is None and run.series["out_dscr"][run.years.index(5)] > 0
    assert run.outputs["out_llcr"] is None


def test_property_tax_follows_hard_costs(bridge, twin_base):
    dense = simulate(TwinParams(site_id="us_ny_jfk", geometry={"kw_per_rack": 30}))
    inputs = workbook_inputs(dense, twin_base, MarketAssumptions(), bridge.manifest)
    hard = sum(inputs[name] for name in CAPEX_KEYS.values())
    base_hard = sum(bridge.manifest["inputs"][name]["value"] for name in CAPEX_KEYS.values())
    assert inputs["tw_property_tax"] == pytest.approx(1_100_000 * hard / base_hard)
    manual = workbook_inputs(dense, twin_base, MarketAssumptions(property_tax_rate=0.01), bridge.manifest)
    assert manual["tw_property_tax"] == pytest.approx(0.01 * hard)


# ------------------------------------------------------------------ règles d'écriture

def test_only_inputs_twin_can_be_written(bridge):
    with pytest.raises(PermissionError):
        bridge.engine.write_inputs({"out_npv": 1.0})
    with pytest.raises(PermissionError):
        bridge.engine.write_inputs({"PUE": 1.4})
    with pytest.raises(PermissionError):  # contrôle côté moteur : Equity_Total est sur Underwriting
        bridge.engine._write({"Equity_Total": 1.0})


def test_bridge_copy_structure(bridge):
    with zipfile.ZipFile(bridge.workbook) as z:
        names = z.namelist()
        inputs_xml = z.read("xl/worksheets/sheet_twin_inputs.xml").decode()
        workbook_xml = z.read("xl/workbook.xml").decode()
    assert "xl/vbaProject.bin" in names and "xl/drawings/drawing1.xml" in names  # macros et boutons conservés
    assert "<f>" not in inputs_xml  # Inputs_Twin ne contient que des constantes
    assert 'iterate="1"' in workbook_xml and 'name="Perm_Debt?"' in workbook_xml
    wb = openpyxl.load_workbook(bridge.workbook)
    uw = wb["Underwriting"]
    for spec in INPUTS:
        if spec.cell:
            assert uw[spec.cell].value == f"={spec.name}", spec.cell
    assert uw["S158"].value == "=Mgmt_Fee_Pct*S148"


def test_bridge_copy_rows_are_well_formed(bridge):
    """Excel refuse d'ouvrir une feuille dont une cellule n'appartient pas à sa ligne <row r=…>."""
    with zipfile.ZipFile(bridge.workbook) as z:
        for name in (n for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")):
            xml = z.read(name).decode()
            previous = 0
            for match in re.finditer(r'<row r="(\d+)"[^>]*?(/>|>(.*?)</row>)', xml, re.S):
                row = int(match.group(1))
                assert row > previous, (name, row)
                previous = row
                for ref in re.findall(r'<c r="[A-Z]+(\d+)"', match.group(3) or ""):
                    assert int(ref) == row, (name, row, ref)


def test_engine_selection():
    if platform.system() not in ("Windows", "Darwin"):
        with pytest.raises(EngineUnavailable):
            select_engine("xlwings")
    with pytest.raises(ValueError):
        select_engine("calc")


# ------------------------------------------------------------------ bout en bout

def test_price_plus_ten_percent_end_to_end(bridge, twin_base):
    base = bridge.run(twin_base, twin_base)
    twin_high = simulate(TwinParams(site_id="us_ny_jfk", price={"flat_usd_per_mwh": 165}))
    high = bridge.run(twin_high, twin_base)
    assert high.outputs["out_opex_energy_y1"] / base.outputs["out_opex_energy_y1"] == pytest.approx(1.10, rel=1e-9)
    assert high.outputs["out_irr_project"] < base.outputs["out_irr_project"]
    assert high.outputs["out_irr_equity"] < base.outputs["out_irr_equity"]
    assert high.recalc_ok and abs(high.ltc_error) <= B.LTC_TOLERANCE


def test_scenario_scaling_and_export(bridge, twin_base, tmp_path):
    twin = simulate(TwinParams(site_id="us_ny_jfk", geometry={"kw_per_rack": 30}))
    run = bridge.run(twin, twin_base, MarketAssumptions(perm_debt=True), export_path=tmp_path / "s.xlsx")
    assert run.outputs["out_it_load_kw"] == pytest.approx(30_000)
    assert run.outputs["out_hard_costs"] == pytest.approx(twin.capex.total_usd)
    assert isinstance(run.outputs["out_dscr_min"], float) and run.outputs["out_llcr"] > 0
    assert run.recalc_ok
    exported = openpyxl.load_workbook(tmp_path / "s.xlsx", data_only=True)
    assert exported["Underwriting"]["E131"].value == pytest.approx(30_000)
    assert exported["Version"]["B2"].value is not None or exported["Version"].max_row > 1  # onglet Version conservé


def test_npv_at_irr_is_zero(bridge):
    """VAN et TRI utilisent la même fenêtre de flux : VAN(TRI) = 0."""
    base = bridge.run_inputs({})
    for irr, npv in (("out_irr_project", "out_npv"), ("out_irr_equity", "out_npv_equity")):
        at_irr = bridge.run_inputs({"tw_discount_rate": base.outputs[irr]})
        assert abs(at_irr.outputs[npv]) < 1e-3 * abs(base.outputs["out_capex_total"]) * 1e-3, npv
